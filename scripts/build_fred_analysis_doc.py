from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION

TARGET = "docs/Freds_Storyboard_Analysis.docx"
INK = "202B3D"
MUTED = "5A6574"
CORAL = "E96C50"
YELLOW = "FFC857"
CREAM = "FFF6E9"
PALE = "F5F7FA"
GREEN = "567F47"


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "247B9A")
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_field(paragraph, field):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, text, end):
        run._r.append(element)


def add_table(doc, headers, rows, widths=None, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        shade(cell, INK)
        set_cell_margins(cell)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.name = "Aptos"
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor(255, 255, 255)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_repeat_table_header(table.rows[0])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cells[i])
            if row_index % 2 == 1:
                shade(cells[i], PALE)
            for paragraph in cells[i].paragraphs:
                paragraph.paragraph_format.space_after = Pt(1)
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(font_size)
                    run.font.color.rgb = RGBColor.from_string(INK)
        if widths:
            for i, width in enumerate(widths):
                cells[i].width = Inches(width)
    if widths:
        for i, width in enumerate(widths):
            table.columns[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_callout(doc, label, body, fill=CREAM, accent=CORAL):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, fill)
    set_cell_margins(cell, top=180, start=210, bottom=180, end=210)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(label.upper())
    r.bold = True
    r.font.name = "Aptos"
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(accent)
    p2 = cell.add_paragraph(body)
    p2.paragraph_format.space_after = Pt(0)
    for r in p2.runs:
        r.font.name = "Aptos"
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor.from_string(INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p


def clear_body(doc):
    body = doc._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


doc = Document(TARGET)
clear_body(doc)
section = doc.sections[0]
section.top_margin = Inches(0.62)
section.bottom_margin = Inches(0.62)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)

normal = doc.styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(9.5)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.08
for style_name, size, color in (("Title", 30, INK), ("Heading 1", 17, INK), ("Heading 2", 12, CORAL), ("Heading 3", 10, GREEN)):
    style = doc.styles[style_name]
    style.font.name = "Aptos Display" if style_name in ("Title", "Heading 1") else "Aptos"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(10 if style_name != "Title" else 0)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.keep_with_next = True

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = header.add_run("FRED'S STORYBOARD LAB  ·  BACKGROUND-FIRST PROTOTYPE")
run.font.name = "Aptos"
run.font.size = Pt(7.5)
run.font.bold = True
run.font.color.rgb = RGBColor.from_string(MUTED)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run("Working brief for Rayvin  ·  20 October 2026  ·  Page ")
run.font.name = "Aptos"
run.font.size = Pt(8)
run.font.color.rgb = RGBColor.from_string(MUTED)
add_field(footer, "PAGE")

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(7)
r = p.add_run("PAID DEMO  /  BUILD BRIEF")
r.bold = True
r.font.name = "Aptos"
r.font.size = Pt(9)
r.font.color.rgb = RGBColor.from_string(CORAL)

p = doc.add_paragraph(style="Title")
p.add_run("Fred's Script-to-\nStoryboard Demo")
p = doc.add_paragraph("Reference analysis  ·  background-first architecture  ·  prototype build plan")
p.paragraph_format.space_after = Pt(14)
for r in p.runs:
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor.from_string(MUTED)

add_callout(doc, "Confirmed product contract", "Paste the full script → AI 1 writes the full scene-by-scene storyboard → AI 2 creates up to 10 distinct, scene-specific background plates per batch. Animators draw the people, body parts, moving props, text, and motion separately.", fill="FFF0E8")

doc.add_heading("1. What Fred asked for", level=1)
doc.add_paragraph("Fred's initial description was a script-in, storyboard-out tool that identifies scenes and creates their backgrounds. His clarification resolves the key scope: the generated images are environment plates only. The storyboard should make an animator's work faster by deciding what each scene needs to look like and where it takes place.")
add_bullet(doc, "Keep one hand-drawn art language throughout the video, but change the setting, framing, and environment to fit each scene.")
add_bullet(doc, "Generate no characters, animals, body parts, hands, loose props, overlay text, or animated effects.")
add_bullet(doc, "The storyboard can include many more scenes than the demo can paint in one batch. Generate no more than 10 distinct backgrounds per batch and leave the remainder visible for later.")

doc.add_heading("2. What the references establish", level=1)
doc.add_paragraph("The two topic-matched videos, scripts, and CSV workbooks are the strongest evidence for this prototype. The channels add audience and packaging context; the CSV columns show how the team translates narration into animator instructions.")
add_table(doc, ["Reference workbook", "Storyboard pattern", "How the planner uses it"], [
    ["Every Weird Body Glitch", "81 scenes through 17:14; median span 11 seconds; long rows include multiple cuts.", "Split long narration into clear visual beats. Preserve menus, rooms, city scenes, chalkboards, gradients, anatomy inserts, and diagrams as distinct settings."],
    ["Strangest Things Your Body Does", "166 scenes through 15:18; median span 5 seconds; frequent short beats.", "Use concise scene rows and let the background follow the beat, including yellow infographic fields, kitchens, streets, rooms, and offices."],
], widths=[1.55, 2.35, 3.3], font_size=8.3)
add_callout(doc, "Style, not one setting", "The yellow living room is a reference for line quality, flat fills, and drawn perspective. It should not become the default scene. A street, classroom, gradient, kitchen, hospital, or field can use the same art style.", fill="F3F8F0", accent=GREEN)

doc.add_heading("3. The storyboard is the main deliverable", level=1)
doc.add_paragraph("Every row ties one exact script excerpt to an action, an environment, and the animator/editor's separate instructions. The first eight CSV columns must stay in Fred's order:")
add_table(doc, ["CSV column", "Use"], [
    ["Scene", "Number every visual beat in the complete script."],
    ["Timestamp", "Copy supplied timecodes; leave blank if none were given."],
    ["Voiceover", "Keep an exact excerpt from the original script."],
    ["What's happening", "Describe one animator-readable beat; split multi-cut rows when needed."],
    ["Background", "Describe only the scene-specific environment or visual field."],
    ["Hand-drawn characters & props", "Keep characters, anatomy, loose props, and moving layers for the animator."],
    ["On-screen text", "Preserve exact overlay wording as metadata."],
    ["Transitions/Motion", "Preserve hard cuts, zooms, fades, and movement as edit notes."],
], widths=[2.1, 5.1], font_size=8.4)
add_callout(doc, "Image input boundary", "The image model receives the Background field plus shared style rules only. It does not receive the action, voiceover, character/prop directions, on-screen text, or transitions.", fill="FFF5D6", accent="9B6B00")

doc.add_heading("4. Shared drawing style and background rules", level=1)
add_bullet(doc, "Bold, slightly irregular black outlines; simple readable shapes; flat fills; sparse detail; restrained shading.")
add_bullet(doc, "Choose the palette for the setting: cream, yellow, ochre, amber, brown, charcoal, green, coral, or blue. Do not force one palette on every scene.")
add_bullet(doc, "Include fixed environmental details only when the Background field calls for them: walls, floors, windows, sofa, shelves, chalkboard, plants, lamps, street, or landscape.")
add_bullet(doc, "Exclude people, faces, animals, silhouettes, body parts, organs, hands, clothing, handheld objects, diagrams, lettering, logos, borders, UI, motion blur, and animation effects.")
add_bullet(doc, "Any Drive reference image containing a person is style evidence only. The image prompt must explicitly tell the model to omit that subject.")

doc.add_heading("5. Prototype architecture", level=1)
add_table(doc, ["Stage", "What happens"], [
    ["1 · Build storyboard", "The UI submits the complete script to one fixed n8n route. The workflow reads the style guide, reference index, and two example sheets; a text-only AI Agent creates the full storyboard."],
    ["2 · Save and export", "n8n validates the scene rows, assigns background IDs, appends all scenes to Google Sheets, builds the CSV, and returns it for display/download."],
    ["3 · Start image batch", "The UI asks the same workflow to generate up to 10 background IDs. n8n acknowledges quickly, then processes images one at a time."],
    ["4 · Save each plate", "OpenRouter GPT Image 2.5 Flare is used only for image generation. Successful image binaries go to private Google Drive; status, file ID, prompt, cost, and errors go to the Sheets ledger."],
    ["5 · Poll and preview", "The UI polls job status. A private n8n image route checks the job/background IDs, downloads one Drive file, and streams it to the page without making Drive files public."],
], widths=[1.55, 5.65], font_size=8.4)
add_callout(doc, "One workflow", "Planning, generation, status, and private image preview routes live in one new n8n workflow with labeled sticky notes. The verified Stage 1 one-image workflow stays intact as a fallback; it is not used by the background-only UI.", fill="F3F6FB", accent=INK)

doc.add_heading("6. Batch notes and visible feedback", level=1)
doc.add_paragraph("The full storyboard is never cut down to 10 rows. The 10-image limit applies only to distinct background plates in one generation request. If more unique backgrounds exist, their storyboard rows stay available and their image status is `deferred`.")
add_callout(doc, "Required UI and CSV note", "This is the full storyboard. More background plates remain for a later batch.", fill="FFF0E8")
doc.add_paragraph("The note appears in the UI and in the CSV's `Generation batch note` column, with the deferred count when known. The user can choose another set of up to 10 deferred IDs in a later batch.")

doc.add_heading("7. External storage and reliability", level=1)
add_bullet(doc, "Google Sheets stores the style guide, reference manifest, job records, and scene/image status rows. Google Drive stores generated image binaries.")
add_bullet(doc, "No n8n Data Tables, n8n Memory, or internal workflow storage are used as the database.")
add_bullet(doc, "Plan and generate routes acknowledge quickly; status polling avoids long browser requests. Image calls are serial and capped at 10 because the Heroku n8n host has a 512 MB memory limit.")
add_bullet(doc, "A failed background is recorded while successful images remain. No silent auto-regeneration loop. The model stays server-side; the browser never gets an API key.")
add_bullet(doc, "CORS restricts browser origins but does not authenticate direct callers. Keep input limits, validate job/background IDs, and avoid publishing private Drive images.")

doc.add_heading("8. Acceptance checks", level=1)
for item in [
    "The full script appears in ordered scene rows and can be exported in Fred's exact eight-column format.",
    "Timecodes are preserved only when supplied; narration excerpts remain exact.",
    "Every Background field fits that scene; the shared art style remains consistent across different settings.",
    "Generated plates contain no characters, animals, anatomy, loose props, text, or motion.",
    "Each batch makes at most 10 distinct image calls, saves each successful binary to Drive, and reports failures/deferred plates.",
    "The note about additional storyboard backgrounds appears in the UI and downloaded CSV.",
    "The UI displays saved image previews via the private image route, and the same workflow creates the CSV.",
]:
    add_bullet(doc, item)

doc.add_heading("Reference ledger", level=1)
sources = [
    ("EverythingProfessor", "https://www.youtube.com/@EverythingProfessor"),
    ("The Historian Professor", "https://www.youtube.com/@TheHistorianProfessor"),
    ("WildlifeProfessor", "https://www.youtube.com/@WildlifeProfessor"),
    ("EverythingProfessor Compilations", "https://www.youtube.com/@EverythingProfessorComps"),
    ("Earlier reference video", "https://drive.google.com/file/d/1eRmvxfh5Ju4cOmTdWt4FSQCsohyM4ISW/view?usp=sharing"),
    ("Every Weird Body Glitch video", "https://www.youtube.com/watch?v=meH-Q1XgLBs"),
    ("Every Weird Body Glitch script", "https://docs.google.com/document/d/19D5RIF-E3ZB-fs_esrgCIw65Mt3jPjBAc2YzREyOvKE/edit?tab=t.0"),
    ("Strangest Things video", "https://www.youtube.com/watch?v=7tPPGl6Q_9I"),
    ("Strangest Things script", "https://docs.google.com/document/d/1vtfXQZaVTlnAxQ4tUsJ2QuKrXiQl7PR4kdYjMc-gsKY/edit?tab=t.0"),
    ("Fred's character/style folder", "https://drive.google.com/drive/folders/1KLTxXSsptR9OguL30DCBLX72GX2rtUBo"),
    ("Imported style and reference index", "https://docs.google.com/spreadsheets/d/1RL7C_3kg7Ua47R7BnQFi4YvsgDyecIQS7yi1uCAT-gg/edit?ouid=112934907610360798465"),
    ("Rayvin's portfolio", "https://rayvingianan.r4yv.tech/"),
    ("Shop Floor Agent repository", "https://github.com/rayvin04/shop-floor-ai-rayvin-gianan"),
    ("OpenRouter GPT Image Flare", "https://openrouter.ai/openai/gpt-image-2.5-flare"),
]
for label, url in sources:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(1)
    add_hyperlink(p, label, url)

doc.save(TARGET)
print(TARGET)
