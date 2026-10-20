from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).resolve().parent.parent / "docs" / "Freds_Storyboard_Analysis.docx"
INK = "202A3B"
MUTED = "647084"
CORAL = "EF674D"
YELLOW = "FFD166"
PALE = "FFF4D8"
BLUE = "EAF6FC"
GREEN = "EAF5E7"
LIGHT = "F5F7FA"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, start=130, bottom=100, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D6DCE4", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_hyperlink(paragraph, text, url, color=CORAL):
    part = paragraph.part
    rid = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    rpr.append(c)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    run.append(rpr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    display = OxmlElement("w:t")
    display.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, display, end])


def set_run(run, size=10.5, color=INK, bold=False, italic=False, font="Aptos"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic
    return run


def add_text(paragraph, text, size=10.5, color=INK, bold=False, italic=False):
    return set_run(paragraph.add_run(text), size, color, bold, italic)


def add_body(doc, text, after=5, before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.12
    add_text(p, text)
    return p


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.left_indent = Inches(0.22 + 0.18 * level)
    p.paragraph_format.first_line_indent = Inches(-0.13)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.06
    add_text(p, text, size=10)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    add_text(p, text, size=17 if level == 1 else 12, color=INK if level == 1 else CORAL, bold=True)
    return p


def add_callout(doc, title, text, fill=PALE, accent=CORAL):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(6.8)
    cell = table.cell(0, 0)
    cell.width = Inches(6.8)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=145, start=190, bottom=135, end=190)
    set_table_borders(table, color=accent, size="9")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    add_text(p, title, size=10, color=accent, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.08
    add_text(p2, text, size=10.2, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def build_table(doc, headers, rows, widths, header_fill=INK, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    for i, width in enumerate(widths):
        table.columns[i].width = Inches(width)
        cell = table.rows[0].cells[i]
        cell.width = Inches(width)
        set_cell_shading(cell, header_fill)
        set_cell_margins(cell, top=90, start=100, bottom=90, end=100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        add_text(p, headers[i], size=8.5, color=WHITE, bold=True)
    set_repeat_table_header(table.rows[0])
    for row_index, row_data in enumerate(rows):
        cells = table.add_row().cells
        fill = WHITE if row_index % 2 == 0 else LIGHT
        for i, value in enumerate(row_data):
            cells[i].width = Inches(widths[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cells[i], fill)
            set_cell_margins(cells[i], top=85, start=100, bottom=85, end=100)
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.02
            add_text(p, value, size=font_size, color=INK, bold=(i == 0))
    set_table_borders(table, color="D7DDE5", size="5")
    return table


def add_flow_table(doc):
    rows = [
        ("01  UI", "Paste a script, select a small scene batch, and review the storyboard."),
        ("02  JOB START", "Validate, create a job ID, save processing state, return a quick acknowledgement."),
        ("03  PLAN + RETRIEVE", "Vercel text model plans beats; style brief and matching references are retrieved per scene."),
        ("04  IMAGE PASS", "OpenRouter GPT Image 2.5 Flare generates one image per scene from text plus selected reference images."),
        ("05  STATUS + REVIEW", "UI polls the same n8n workflow, previews images, captures ratings and comments."),
        ("06  DELIVERY", "Upload the completed batch to Drive, then send one Slack notice to the editing team."),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    table.autofit = False
    table.columns[0].width = Inches(1.52)
    table.columns[1].width = Inches(5.28)
    for i, (label, desc) in enumerate(rows):
        left, right = table.rows[i].cells
        left.width = Inches(1.52)
        right.width = Inches(5.28)
        set_cell_margins(left, top=88, start=115, bottom=88, end=115)
        set_cell_margins(right, top=88, start=125, bottom=88, end=125)
        set_cell_shading(left, [YELLOW, BLUE, PALE, GREEN, "F1EAFB", YELLOW][i])
        set_cell_shading(right, WHITE if i % 2 == 0 else LIGHT)
        p = left.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        add_text(p, label, size=8.5, bold=True)
        p = right.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        add_text(p, desc, size=9.2)
    set_table_borders(table, color="CBD3DD", size="5")
    return table


def add_source(doc, label, url, note=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Inches(0.05)
    add_text(p, label, size=9.2, bold=True)
    add_text(p, "  ")
    add_hyperlink(p, "Open reference", url)
    if note:
        add_text(p, "  " + note, size=8.7, color=MUTED)
    return p


def configure(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.32)
    section.footer_distance = Inches(0.32)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.1
    for name, size, color in (("Heading 1", 17, INK), ("Heading 2", 12, CORAL), ("Heading 3", 10.5, INK)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(header, "FRED  /  SCRIPT → STORYBOARD", size=8, color=MUTED, bold=True)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(footer, "REFERENCE ANALYSIS  •  20 OCT 2026     |     ", size=8, color=MUTED)
    add_page_field(footer)


def build_report():
    doc = Document()
    configure(doc)

    # Page 1: executive brief
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(22)
    p.paragraph_format.space_after = Pt(4)
    add_text(p, "PAID DEMO  /  BUILD BRIEF", size=9, color=CORAL, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    add_text(p, "Fred's Script-to-\nStoryboard Demo", size=31, color=INK, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    add_text(p, "Reference analysis • prototype architecture • phased build plan", size=12, color=MUTED)

    add_callout(
        doc,
        "THE FIRST USEFUL RESULT",
        "A user pastes a script, receives storyboard rows in Fred's CSV format, and gets one complete still frame for each selected beat. The user reviews and downloads the small batch; timing, overlay text, and edit cues remain attached as scene metadata.",
        fill=PALE,
    )
    add_heading(doc, "What Fred is asking for", 1)
    add_body(
        doc,
        "Fred wants a practical script-to-storyboard automation: the script is segmented into rows that describe narration, action, background, drawn characters/props, optional on-screen text, and transitions/motion. The image model then creates a matching full scene frame for each selected row in a consistent art style. The aim is to reduce repeated manual visual planning.",
    )
    add_heading(doc, "Recommendation", 2)
    add_bullet(doc, "Keep a single n8n workflow behind the reviewer-facing UI; expose fixed start, status, feedback, and delivery routes in the UI code.")
    add_bullet(doc, "Use the two new body videos, their matching scripts, and the supplied storyboard CSVs as the direct format/content references. Keep the older asteroid sample as broader pacing context only.")
    add_bullet(doc, "Keep the first run small: 1–3 scenes, one image per scene, manual approval before full-script batches.")

    add_callout(
        doc,
        "REFERENCE PRIORITY",
        "The new Every Weird Body Glitch and Strangest Things Your Body Does to Keep You Alive videos match the two script/CSV references. Follow the CSV row structure for image planning. The earlier asteroid sample has a different topic and is useful only for broader pacing/style context.",
        fill=BLUE,
        accent="3C7291",
    )
    add_body(doc, "Prepared for Rayvin to use as a build reference and to brief Fred. New script and CSV references reviewed 20 October 2026.", after=0)

    # Page 2: editorial and script references
    doc.add_page_break()
    add_heading(doc, "1. Channel and script analysis", 1)
    add_body(doc, "The channels share an accessible explainer promise: take a broad or surprising topic and make it easy to follow through a clear hook, short sections, and visual changes that support the narration.")
    build_table(
        doc,
        ["Reference", "Observed format", "Use in this demo"],
        [
            ("EverythingProfessor", "“Complex things, simply explained”; list-style psychology, health, sleep, and general explainers.", "Audience packaging and pacing context."),
            ("Professor Historian", "Entertaining history explainers; list, mystery, and first-person framing.", "Shows the broader explainer brand."),
            ("WildlifeProfessor", "Animal behavior, capability, and survival topics.", "Adjacent channel context, not the first content."),
            ("EverythingProfessor Compilations", "Long compilations assembled around topic clusters.", "Potential later use for repeatable scene assets."),
        ],
        [1.55, 2.75, 2.5],
        font_size=8.7,
    )
    add_heading(doc, "Supplied script: 15 topic beats", 2)
    topics = [
        "Phantom phone vibration", "Pins and needles", "Eye twitching", "Brain zaps", "Ear ringing / tinnitus",
        "Skin crawling", "Unexplained goosebumps", "Random muscle twitching", "Dizziness on standing",
        "Stomach-drop feeling", "Lump-in-throat feeling", "Sharp chest / rib stitch", "Face flushing",
        "Heart palpitations", "Internal vibration sensation",
    ]
    for start in range(0, len(topics), 3):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.left_indent = Inches(0.12)
        add_text(p, "   •   ".join(topics[start:start + 3]), size=9.2)
    add_callout(
        doc,
        "SCRIPT VOICE",
        "Casual and wry, moving from a familiar sensation to a plain-language explanation, a relatable analogy, and a practical caveat. The image system should explain each beat without inventing medical facts or making it look like a diagnosis.",
        fill=GREEN,
        accent="4B8250",
    )
    add_body(doc, "An MP3 named audio_159655.mp3 is present in the script folder. I have not verified its timecodes or produced a transcript, so the first storyboard should segment from the supplied text only.", after=0)

    # Page 3: storyboard CSV contract
    doc.add_page_break()
    add_heading(doc, "2. How the storyboard rows work", 1)
    add_body(doc, "The CSV is the planner's scene format. It ties a narration excerpt to a visual action, setting, drawn subjects/props, optional overlay text, and editor cues. The image request uses the content fields to create one complete frame; it does not render the editing metadata.")
    build_table(
        doc,
        ["CSV column", "Scene field", "How the planner uses it"],
        [
            ("Scene", "sceneNumber / id", "Preserve order and a stable scene identifier."),
            ("Timestamp", "timestamp", "Copy exact mm:ss-mm:ss values; leave blank when only a script is supplied."),
            ("Voiceover", "narration", "Keep an exact concise excerpt from the script."),
            ("What's happening", "whatHappening", "Choose one clear still moment; split multi-cut rows when needed."),
            ("Background", "background", "Describe the setting or diagram canvas."),
            ("Hand-drawn characters & props", "handDrawnElements", "List the visible characters, anatomy, and props."),
            ("On-screen text", "onScreenText", "Preserve exact overlay text as metadata; do not draw it."),
            ("Transitions/Motion", "transitionsMotion", "Keep editor cues as metadata; freeze any depicted action."),
        ],
        [1.85, 1.55, 3.4],
        font_size=8.1,
    )
    build_table(
        doc,
        ["Workbook", "Rows / timeline", "Timing notes"],
        [
            ("Every Weird Body Glitch", "81 rows; ends at 17:14", "Median 11s; mean 12.7s; max 39s; one zero-duration row and five gaps; no overlaps."),
            ("Strangest Things Your Body Does", "166 rows; ends at 15:18", "Median 5s; max 14s; no zero-duration rows, gaps, or overlaps."),
        ],
        [1.8, 1.55, 3.45],
        font_size=8.1,
    )
    add_heading(doc, "Recurring visual families", 2)
    build_table(
        doc,
        ["Family", "CSV examples / image treatment"],
        [
            ("topic_menu", "White field, circular topic icons; reserve space for final title/labels."),
            ("character_environment", "Green-shirt protagonist in the yellow-wall / green-couch living room; build the specified place and props."),
            ("host_lecture", "Professor explaining one idea beside a chalkboard or yellow teaching field."),
            ("single_action_vignette", "One pose, such as reaching into a pocket, pulling a hand from a stove, sneezing, or fainting."),
            ("anatomy_insert", "Clear, simplified, non-graphic body or organ detail."),
            ("concept_metaphor", "One concrete analogy, such as a brain running on a hamster wheel."),
            ("diagram_comparison", "Split screen, body mechanism, cutaway, or selected arrows; labels remain overlay metadata."),
        ],
        [1.8, 5.0],
        font_size=8.0,
    )
    add_callout(
        doc,
        "ONE IMAGE PER VISUAL BEAT",
        "Some CSV rows contain several cuts or actions, especially the 31-second and 37-second examples. Split them into one frozen action per image when the scene budget allows. Each generated image includes the described background plus the listed characters/props; it is not an empty background plate. If the user pasted only script text, do not guess timestamps, overlay labels, or transitions.",
        fill=PALE,
    )

    # Page 4: visual style
    doc.add_page_break()
    add_heading(doc, "3. Visual language and assets", 1)
    add_body(doc, "The supplied references show several connected scene families that share heavy, hand-drawn outlines. The two CSVs alternate topic menus, character vignettes, professor-led explanations, analogies, diagrams, and anatomy inserts. Retrieve the family that matches each scene instead of averaging them into one generic cartoon.")
    build_table(
        doc,
        ["Cue", "Observed reference trait", "Prompting implication"],
        [
            ("Characters", "Rounded heads, minimal eyes and mouths, stick-like limbs, familiar outfits.", "Use one approved character anchor repeatedly."),
            ("Linework", "Heavy, slightly irregular black outlines, especially on the host and anatomy.", "Preserve bold hand-drawn contours; avoid thin corporate vector strokes."),
            ("Color", "Cream faces; warm amber, ochre, brown and charcoal settings; coral/orange organs; occasional blue daylight.", "Keep a warm palette and use lighting to separate the focal subject."),
            ("Composition", "Simple host against layered rooms; wide, medium and oversized close-up shots; isolated organ inserts.", "Select character/environment or anatomy references based on the beat."),
            ("Format", "Sample video is 16:9 and includes editorial overlays.", "Generate clean 16:9 art without titles, captions, logos, or UI."),
        ],
        [1.05, 2.85, 2.9],
        font_size=8.8,
    )
    add_heading(doc, "Style prompt foundation", 2)
    add_callout(
        doc,
        "LOCKED STYLE FOUNDATION",
        "Match the approved Fred references: bold irregular black ink, a pale cream oval host head with sparse facial marks and the same glasses/clothing silhouette where applicable, warm amber-brown environments with layered props and controlled shadows, and simplified coral-orange anatomy inserts with clear dark contours. Choose the matching visual family for each beat. Create one 16:9 educational frame with one focal action. No captions, labels, channel marks, watermarks, UI, photorealism, glossy 3D, anime, or unrelated clip-art.",
        fill=PALE,
    )
    add_heading(doc, "Asset handling", 2)
    add_bullet(doc, "The shared folder contains 18 PNGs, including alternate women, a soldier, lifeguard, handyman, agents, businessmen, a professor, and a living-room scene.")
    add_bullet(doc, "Six added screenshots show distinct character/environment and anatomy examples. Curate separate reference groups; do not average cafe, living-room, bar, and organ-close-up frames into one scene template.")
    add_bullet(doc, "Choose one approved host anchor plus one or two matching scene references. Send the actual image references to the image model when the selected endpoint supports them; text-only style rules are not enough for character identity.")
    add_bullet(doc, "The local prototype's SVG frames are newly drawn placeholders. They are not copied from Fred's character files and are not AI generations.")

    add_heading(doc, "UI direction", 2)
    add_body(doc, "The Shop Floor Agent screenshot shows Rayvin's useful “open a demo and try it now” pattern: a clear client label, short explanation, and one prominent interaction. The second screenshot suggests a playful orange/yellow cartoon mood. This UI combines that easy-to-try flow with an off-white canvas, ink outlines, coral/yellow/sky accents, and a storyboard grid.")
    add_body(doc, "Rayvin's portfolio and the Shop Floor Agent repository show a workflow-first approach: make a web UI for immediate use, keep n8n behind it, ground answers in structured data, and document the workflow. The same pattern fits this demo.", after=0)

    # Page 5: architecture and current n8n state
    doc.add_page_break()
    add_heading(doc, "4. First prototype architecture", 1)
    add_body(doc, "The published path remains synchronous. Draft v2.1 now follows the CSV schema and includes retrieved CSV examples. Manual execution 2021 failed at request-body parsing before any image call; after fixing the expression, execution 2022 generated one frame and returned the storyboard fields with a PNG binary. The provider cost for execution 2022 was not retrieved. The Heroku host returned HTTP 503 after that successful run, so the new draft is not published and the production UI still uses the older active version.")
    add_flow_table(doc)
    add_heading(doc, "Current generation path and planned additions", 2)
    build_table(
        doc,
        ["Stage", "Node / responsibility"],
        [
            ("Input", "Webhook POST accepts script, scene limit, and aspect-ratio label at /webhook/script-to-storyboard."),
            ("Guardrail", "Code rejects empty or over-30,000-character scripts and clamps the scene count to 1–3; it creates an internal request identifier."),
            ("Job control", "The current path is synchronous. The planned route writes a processing job, responds 202, and keeps planning/generation in the same workflow."),
            ("Text model", "openai/gpt-5-mini creates CSV-inspired scene rows with exact narration, action, setting, hand-drawn subjects/props, overlay text, edit cues, and image prompt. Text planning does not call the image model."),
            ("Current style retrieval", "Fred Style Reference Search retrieves matching family rules and selected examples from the two CSVs using keyword overlap. It is not semantic/vector search and does not fetch Drive image files."),
            ("Image generation", "Sequential HTTP Request to OpenRouter POST /api/v1/images using openai/gpt-image-2.5-flare at low quality, one image per selected scene. Prompt contains setting plus drawn subjects/props; time, exact overlay text, and edit cues stay metadata."),
            ("Visible image output", "Convert Scene Image to Binary Preview creates a PNG binary. Restore Scene JSON with Binary carries the CSV-style scene fields and image together. Execution 2022 exposed a 1.25 MB PNG under the node's Binary tab."),
            ("Planned async/status", "Upsert the processing/final job in the jobs Data Table; respond 202 immediately; a status route returns completed, partial, or failed results to UI polling."),
            ("Planned feedback/delivery", "Store per-scene likes and comments; send completed images to the supplied Drive folder; then post one completion notice to Slack #social."),
            ("Failure handling", "Keep provider and job failure routes inside this one workflow, as Rayvin selected."),
        ],
        [1.4, 5.4],
        font_size=8.9,
    )
    add_callout(
        doc,
        "CURRENT DRAFT STATUS AND LIMITS",
        "Manual execution 2022 confirmed one image completed and the scene JSON survived the binary preview path. Its provider-reported usage cost was not retrieved; earlier run 2014 reported $0.00575 for one low-quality image. The latest v2.1 draft remains unpublished because the Heroku host returned HTTP 503 after the run. The active production UI therefore still uses the older version. Async job/status routes, feedback persistence, Drive/Slack delivery, and actual Drive image-reference conditioning remain future work. Drive upload access still needs a real test.",
        fill=BLUE,
        accent="3C7291",
    )
    add_body(doc, "When n8n returns, open execution 2022 and inspect `Restore Scene JSON with Binary` → Binary → data. Then publish the v2.1 draft and check one image through the UI before increasing the scene count.", after=0)

    # Page 6: feedback and reliability
    doc.add_page_break()
    add_heading(doc, "5. Feedback, delivery, and reliability", 1)
    add_heading(doc, "Feedback is retrieval, not automatic training", 2)
    add_body(doc, "A thumbs-up/down and a short comment are useful if stored with the exact scene prompt, style version, model, and any edit request. Later, retrieve a few approved examples for similar scenes and use their prompt patterns as context. This is feedback-guided retrieval; the model weights do not update after each comment.")
    build_table(
        doc,
        ["Prototype", "Next iteration", "Production-ready"],
        [
            ("Always-on style brief plus a small tagged reference manifest; retrieve approved images per scene family.", "Add semantic vector search when tags stop scaling; keep Drive files as image binaries and vector metadata separate.", "Versioned reference curation, prompt/audit trail, and approvals before updating the canon."),
            ("Local frame downloads and JSON scene-plan export.", "Google Drive folder export for editors.", "Job records, durable image objects, idempotent retry, resumable batches."),
            ("Per-scene thumbs/comments with browser fallback.", "Persist to the existing n8n feedback table and retrieve approved examples.", "Quality review gates, rate limits, spend caps, authenticated webhooks, monitoring."),
        ],
        [2.05, 2.35, 2.4],
        font_size=8.4,
    )
    add_heading(doc, "Spend and safety controls", 2)
    add_bullet(doc, "Default to 1–3 images per first run. Do not include an automatic re-generation loop.")
    add_bullet(doc, "Use the image model only for image generation. Scene planning, reference metadata, and feedback summaries stay on the text-model route or deterministic code.")
    add_bullet(doc, "Execution 2014 reported $0.00575 for one low-quality image-generation call. The newer successful execution 2022 did not expose its provider cost in the retrieved result. Inspect provider usage before increasing the scene cap.")
    add_bullet(doc, "For public deployment, authenticate the webhook, restrict CORS to the UI domain, keep credentials server-side, and set a spend limit. Never put API keys in front-end code.")
    add_bullet(doc, "For the health script, keep narration facts and caveats unchanged; use neutral visual metaphors and no diagnostic text.")
    add_callout(
        doc,
        "ERROR AND COMPLETION BEHAVIOR",
        "Rayvin selected same-workflow error handling. The planned delivery route uploads the generated batch to Drive folder 153VPM-IBR1T5UIq-TthqtpklyAGwi-UU and then posts one message to Slack #social using the saved credential. The Drive credential's upload access remains unverified; only the final UI end-to-end test can confirm delivery.",
        fill=GREEN,
        accent="4B8250",
    )

    # Page 7: phases, estimate, sources
    doc.add_page_break()
    add_heading(doc, "6. Phased build and estimate", 1)
    build_table(
        doc,
        ["Phase", "Outcome", "Exit check"],
        [
            ("1. Working demo", "UI, async job/status in one workflow, scene planning, low-cost image call, binary execution preview, per-scene downloads/review.", "One scene reaches both the UI and the n8n execution Binary tab."),
            ("2. Useful demo", "Curated style references, feedback persistence/retrieval, Drive upload, Slack completion notice, error paths.", "An editor can review a complete delivered batch and its feedback is saved."),
            ("3. Production", "Semantic vector retrieval if needed, access controls, spend limits, durable object storage, recovery, monitoring, runbook.", "Outputs are traceable to style version, refs, prompt, and reviewer feedback."),
        ],
        [1.25, 3.25, 2.3],
        font_size=8.6,
    )
    add_heading(doc, "Time estimate for Fred", 2)
    add_callout(
        doc,
        "QUOTE-READY ESTIMATE",
        "“For a focused demo, I estimate 2–3 working days once the reference files and image provider are accessible: a small web UI, script-to-scene planning, up to three generated frames, downloads, and a review step. A polished team-ready version with Drive delivery, persistent feedback, retries, notifications, and deployment hardening would likely take another 3–5 working days.”",
        fill=PALE,
    )
    add_body(doc, "Text planning uses the saved Vercel credential and image generation uses the saved OpenRouter credential. The output Drive folder is selected, but n8n upload permission needs verification. The n8n host outage is delaying implementation, and final timing depends on service recovery and endpoint reference-image support.", after=5)
    add_heading(doc, "Open decisions after the prototype", 2)
    add_bullet(doc, "Which one character is the primary anchor for the body-glitch video?")
    add_bullet(doc, "Confirm with Fred whether the editing team expects complete scene frames or separate empty background plates; the demo currently treats each CSV row as a complete frame with background plus drawn subjects/props.")
    add_bullet(doc, "Which specific character image should be the canonical host anchor, and should it appear in every visual family?")
    add_bullet(doc, "Confirm n8n's saved Google Drive credential can upload to the supplied `Demo Test Upload` folder.")
    add_bullet(doc, "Confirm #social is the right editor notification channel before sending the first live delivery alert.")

    add_heading(doc, "Sources checked", 2)
    add_source(doc, "EverythingProfessor", "https://www.youtube.com/@EverythingProfessor")
    add_source(doc, "Professor Historian", "https://www.youtube.com/@TheHistorianProfessor")
    add_source(doc, "WildlifeProfessor", "https://www.youtube.com/@WildlifeProfessor")
    add_source(doc, "EverythingProfessor Compilations", "https://www.youtube.com/@EverythingProfessorComps")
    add_source(doc, "Every Weird Body Glitch video", "https://www.youtube.com/watch?v=meH-Q1XgLBs")
    add_source(doc, "Every Weird Body Glitch script", "https://docs.google.com/document/d/19D5RIF-E3ZB-fs_esrgCIw65Mt3jPjBAc2YzREyOvKE/edit?tab=t.0")
    add_source(doc, "Strangest Things Your Body Does video", "https://www.youtube.com/watch?v=7tPPGl6Q_9I")
    add_source(doc, "Strangest Things Your Body Does script", "https://docs.google.com/document/d/1vtfXQZaVTlnAxQ4tUsJ2QuKrXiQl7PR4kdYjMc-gsKY/edit?tab=t.0#heading=h.gzdctbyeq5dv")
    add_source(doc, "Earlier sample video (asteroid topic)", "https://drive.google.com/file/d/1eRmvxfh5Ju4cOmTdWt4FSQCsohyM4ISW/view?usp=sharing", "Broader visual/pacing context")
    add_source(doc, "Script / voiceover folder", "https://drive.google.com/drive/u/1/folders/1q_BkGppDA4k189FXZxXewRWi9yJF_r4H")
    add_source(doc, "Character / style folder", "https://drive.google.com/drive/folders/1KLTxXSsptR9OguL30DCBLX72GX2rtUBo?usp=sharing")
    add_source(doc, "Portfolio and Shop Floor Agent", "https://rayvingianan.r4yv.tech/", "Repo: github.com/rayvin04/shop-floor-ai-rayvin-gianan")
    add_source(doc, "OpenRouter GPT Image 2.5 Flare listing", "https://openrouter.ai/openai/gpt-image-2.5-flare", "Image generation; inspect endpoint-specific reference support")
    add_source(doc, "OpenRouter Images API guide", "https://openrouter.ai/blog/tutorials/image-generation/", "Includes input_references")
    add_body(doc, "The two uploaded workbooks analyzed were 01_every_weird_body_glitch.csv.xlsx and 03_strangest_things_your_body_does.csv.xlsx. Both use the eight-column scene format summarized in section 2.", after=0)

    doc.core_properties.title = "Fred's Script-to-Storyboard Demo — Reference Analysis"
    doc.core_properties.subject = "Reference analysis, architecture, and phased build estimate"
    doc.core_properties.author = "Rayvin Gianan"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_report()
